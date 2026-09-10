import hashlib
import io
import json
import uuid
from typing import Any

import numpy as np
import streamlit as st
from docx import Document
from openai import OpenAI
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


st.set_page_config(
    page_title="Linear LLM Workflow Builder",
    page_icon="🔗",
    layout="wide",
)

st.markdown(
    """
    <style>
      .block-container {max-width: 1280px; padding-top: 1.8rem; padding-bottom: 3rem;}
      .small-note {color: #6b7280; font-size: 0.88rem;}
      .workflow-box {
          border: 1px solid rgba(128,128,128,.25);
          border-radius: 12px;
          padding: 12px 14px;
          margin-bottom: 8px;
      }
    </style>
    """,
    unsafe_allow_html=True,
)

EMBEDDING_MODEL = "text-embedding-3-small"
DEFAULT_MODEL = "gpt-5-mini"
CHUNK_SIZE = 1400
CHUNK_OVERLAP = 200
MAX_RAG_CONTEXT_CHARS = 14000
MAX_EMBED_QUERY_CHARS = 6000


def make_agent(index: int) -> dict[str, Any]:
    return {
        "id": f"agent_{uuid.uuid4().hex[:10]}",
        "name": f"Agent {index}",
        "model": DEFAULT_MODEL,
        "system_prompt": "당신은 맡은 역할을 정확하게 수행하는 전문 LLM 에이전트입니다.",
        "extra_prompt": "",
        "rag_enabled": False,
        "top_k": 4,
        "max_output_tokens": 2000,
        "chatbot_mode": False,
    }


def init_state() -> None:
    if "agents" not in st.session_state:
        st.session_state.agents = [make_agent(1), make_agent(2)]
    if "rag_cache" not in st.session_state:
        st.session_state.rag_cache = {}
    if "last_results" not in st.session_state:
        st.session_state.last_results = []
    if "api_key" not in st.session_state:
        st.session_state.api_key = ""
    if "chat_messages" not in st.session_state:
        st.session_state.chat_messages = []
    if "chat_context" not in st.session_state:
        st.session_state.chat_context = ""
    if "chat_agent_id" not in st.session_state:
        st.session_state.chat_agent_id = None


def read_uploaded_file(uploaded_file) -> str:
    data = uploaded_file.getvalue()
    suffix = uploaded_file.name.rsplit(".", 1)[-1].lower() if "." in uploaded_file.name else ""

    if suffix == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()

    if suffix == "docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts).strip()

    if suffix == "pptx":
        presentation = Presentation(io.BytesIO(data))
        parts = []
        for slide_no, slide in enumerate(presentation.slides, start=1):
            slide_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_parts.append(shape.text.strip())
            if slide_parts:
                parts.append(f"[Slide {slide_no}]\n" + "\n".join(slide_parts))
        return "\n\n".join(parts).strip()

    if suffix == "xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    parts.append(" | ".join(values))
        return "\n".join(parts).strip()

    for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            pass

    return data.decode("utf-8", errors="ignore").strip()


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []

    chunks = []
    start = 0
    length = len(text)

    while start < length:
        end = min(start + chunk_size, length)
        if end < length:
            candidates = [
                text.rfind("\n\n", start, end),
                text.rfind("\n", start, end),
                text.rfind(". ", start, end),
            ]
            split_at = max(candidates)
            if split_at > start + chunk_size // 2:
                end = split_at + 1

        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)

        if end >= length:
            break
        start = max(end - overlap, start + 1)

    return chunks


def files_digest(files) -> str:
    h = hashlib.sha256()
    for f in sorted(files, key=lambda x: x.name):
        h.update(f.name.encode("utf-8", errors="ignore"))
        h.update(f.getvalue())
    return h.hexdigest()


def embed_texts(client: OpenAI, texts: list[str], batch_size: int = 96) -> np.ndarray:
    vectors = []
    for start in range(0, len(texts), batch_size):
        batch = texts[start : start + batch_size]
        response = client.embeddings.create(
            model=EMBEDDING_MODEL,
            input=batch,
        )
        vectors.extend(item.embedding for item in response.data)

    matrix = np.asarray(vectors, dtype=np.float32)
    norms = np.linalg.norm(matrix, axis=1, keepdims=True)
    return matrix / np.clip(norms, 1e-12, None)


def build_or_get_rag_index(client: OpenAI, agent_id: str, files) -> dict[str, Any]:
    digest = files_digest(files)
    cache_key = f"{digest}:{EMBEDDING_MODEL}:{CHUNK_SIZE}:{CHUNK_OVERLAP}"
    cached = st.session_state.rag_cache.get(agent_id)

    if cached and cached.get("cache_key") == cache_key:
        return cached

    chunk_records = []
    extraction_errors = []

    for uploaded in files:
        try:
            text = read_uploaded_file(uploaded)
            if not text:
                extraction_errors.append(f"{uploaded.name}: 추출 가능한 텍스트가 없습니다.")
                continue

            for idx, chunk in enumerate(chunk_text(text), start=1):
                chunk_records.append(
                    {
                        "source": uploaded.name,
                        "chunk_no": idx,
                        "text": chunk,
                    }
                )
        except Exception as exc:
            extraction_errors.append(f"{uploaded.name}: {exc}")

    if not chunk_records:
        detail = "\n".join(extraction_errors) if extraction_errors else "업로드된 파일에서 텍스트를 찾지 못했습니다."
        raise ValueError(detail)

    vectors = embed_texts(client, [item["text"] for item in chunk_records])
    index = {
        "cache_key": cache_key,
        "chunks": chunk_records,
        "vectors": vectors,
        "errors": extraction_errors,
    }
    st.session_state.rag_cache[agent_id] = index
    return index


def retrieve_context(
    client: OpenAI,
    agent_id: str,
    files,
    query: str,
    top_k: int,
) -> tuple[str, list[dict[str, Any]], list[str]]:
    index = build_or_get_rag_index(client, agent_id, files)
    safe_query = query.strip()
    if len(safe_query) > MAX_EMBED_QUERY_CHARS:
        # Embedding models have an input-token limit. Keep the beginning and end
        # so both the main task and recent downstream instructions remain represented.
        half = MAX_EMBED_QUERY_CHARS // 2
        safe_query = safe_query[:half] + "\n...\n" + safe_query[-half:]
    query_vec = embed_texts(client, [safe_query])[0]
    scores = index["vectors"] @ query_vec

    top_k = min(top_k, len(index["chunks"]))
    order = np.argsort(scores)[::-1][:top_k]

    selected = []
    context_parts = []
    char_count = 0

    for rank, idx in enumerate(order, start=1):
        record = index["chunks"][int(idx)]
        score = float(scores[int(idx)])
        block = (
            f"[Reference {rank} | source={record['source']} | chunk={record['chunk_no']}]\n"
            f"{record['text']}"
        )

        if context_parts and char_count + len(block) > MAX_RAG_CONTEXT_CHARS:
            break

        context_parts.append(block)
        char_count += len(block)
        selected.append(
            {
                "rank": rank,
                "source": record["source"],
                "chunk_no": record["chunk_no"],
                "score": score,
                "text": record["text"],
            }
        )

    return "\n\n".join(context_parts), selected, index.get("errors", [])


def compose_user_input(workflow_input: str, extra_prompt: str, rag_context: str) -> str:
    parts = [
        "## WORKFLOW INPUT",
        workflow_input.strip(),
    ]

    if extra_prompt.strip():
        parts.extend(
            [
                "",
                "## ADDITIONAL INSTRUCTION FOR THIS AGENT",
                extra_prompt.strip(),
            ]
        )

    if rag_context.strip():
        parts.extend(
            [
                "",
                "## RETRIEVED REFERENCE CONTEXT",
                rag_context.strip(),
                "",
                "Use the reference context only when it is relevant. "
                "If the references do not support a factual claim, do not invent support.",
            ]
        )

    return "\n".join(parts)


def call_agent(
    client: OpenAI,
    agent: dict[str, Any],
    workflow_input: str,
    rag_context: str,
) -> tuple[str, str]:
    user_input = compose_user_input(
        workflow_input=workflow_input,
        extra_prompt=agent["extra_prompt"],
        rag_context=rag_context,
    )

    response = client.responses.create(
        model=agent["model"].strip(),
        instructions=agent["system_prompt"].strip(),
        input=user_input,
        max_output_tokens=int(agent["max_output_tokens"]),
        store=False,
    )

    output_text = (response.output_text or "").strip()
    if not output_text:
        raise RuntimeError("모델이 텍스트 출력을 반환하지 않았습니다.")

    return output_text, user_input


def call_chat_agent(
    client: OpenAI,
    agent: dict[str, Any],
    workflow_memory: str,
    chat_messages: list[dict[str, str]],
    user_prompt: str,
    rag_context: str = "",
) -> str:
    instructions = agent["system_prompt"].strip()

    if agent["extra_prompt"].strip():
        instructions += (
            "\n\n## ADDITIONAL INSTRUCTION\n"
            + agent["extra_prompt"].strip()
        )

    if workflow_memory.strip():
        instructions += (
            "\n\n## WORKFLOW MEMORY\n"
            "아래 내용은 이 대화가 시작되기 전에 완료된 linear workflow의 결과입니다. "
            "후속 대화에서 이 내용을 기억하고 필요한 경우 근거로 활용하세요.\n\n"
            + workflow_memory.strip()
        )

    current_user_input = user_prompt.strip()
    if rag_context.strip():
        current_user_input += (
            "\n\n## RETRIEVED REFERENCE CONTEXT\n"
            + rag_context.strip()
            + "\n\n참조 문맥이 관련 있을 때만 활용하고, 근거가 없는 내용은 지어내지 마세요."
        )

    conversation = [
        {"role": message["role"], "content": message["content"]}
        for message in chat_messages
    ]
    conversation.append({"role": "user", "content": current_user_input})

    response = client.responses.create(
        model=agent["model"].strip(),
        instructions=instructions,
        input=conversation,
        max_output_tokens=int(agent["max_output_tokens"]),
        store=False,
    )

    output_text = (response.output_text or "").strip()
    if not output_text:
        raise RuntimeError("챗봇이 텍스트 출력을 반환하지 않았습니다.")

    return output_text


def workflow_config_json() -> str:
    payload = {
        "version": 1,
        "type": "linear_llm_workflow",
        "agents": [
            {
                "name": a["name"],
                "model": a["model"],
                "system_prompt": a["system_prompt"],
                "extra_prompt": a["extra_prompt"],
                "rag_enabled": a["rag_enabled"],
                "top_k": a["top_k"],
                "max_output_tokens": a["max_output_tokens"],
                "chatbot_mode": a.get("chatbot_mode", False),
            }
            for a in st.session_state.agents
        ],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2)


init_state()

st.title("🔗 Linear LLM Workflow Builder")
st.caption(
    "여러 LLM Agent를 순서대로 연결하고, 앞 Agent의 출력을 다음 Agent의 입력으로 전달합니다. "
    "각 Agent에는 독립적인 System Prompt, 추가 지시문, 선택적 RAG를 설정할 수 있습니다."
)

with st.sidebar:
    st.header("OpenAI API")
    st.text_input(
        "API Key",
        key="api_key",
        type="password",
        placeholder="sk-...",
        help="키는 코드에 저장하지 않고 현재 Streamlit 세션에서만 사용합니다.",
    )
    st.caption("Community Cloud에 공개 배포할 경우, 각 사용자가 자신의 API Key를 입력하는 방식입니다.")

    st.divider()
    st.subheader("Workflow")
    c1, c2 = st.columns(2)

    with c1:
        if st.button("➕ Agent", use_container_width=True):
            st.session_state.agents.append(make_agent(len(st.session_state.agents) + 1))
            st.rerun()

    with c2:
        if st.button("🧹 Reset", use_container_width=True):
            st.session_state.agents = [make_agent(1), make_agent(2)]
            st.session_state.rag_cache = {}
            st.session_state.last_results = []
            st.session_state.chat_messages = []
            st.session_state.chat_context = ""
            st.session_state.chat_agent_id = None
            st.rerun()

    st.download_button(
        "⬇️ Workflow 설정 JSON",
        data=workflow_config_json(),
        file_name="linear_workflow_config.json",
        mime="application/json",
        use_container_width=True,
    )

    st.divider()
    st.caption(
        "지원 RAG 파일: PDF, DOCX, PPTX, XLSX, TXT, MD, CSV, JSON. "
        "스캔 이미지형 PDF는 텍스트가 추출되지 않을 수 있습니다."
    )


if st.session_state.agents:
    flow_names = "  →  ".join(
        f"**{i + 1}. {agent['name']}**"
        + (" 💬" if i == len(st.session_state.agents) - 1 and agent.get("chatbot_mode", False) else "")
        for i, agent in enumerate(st.session_state.agents)
    )
    st.markdown(f"### Workflow\n{flow_names}")
else:
    st.info("왼쪽 사이드바의 **Agent 추가** 버튼으로 첫 Agent를 만들어 주세요.")


agent_uploads: dict[str, Any] = {}

st.markdown("### Agent 편집")

for i, agent in enumerate(st.session_state.agents):
    agent_id = agent["id"]
    with st.expander(f"{i + 1}. {agent['name']}", expanded=True):
        top = st.columns([4, 1, 1, 1])

        with top[0]:
            agent["name"] = st.text_input(
                "Agent 이름",
                value=agent["name"],
                key=f"name_{agent_id}",
            )

        with top[1]:
            if st.button("⬆️", key=f"up_{agent_id}", disabled=i == 0, use_container_width=True):
                st.session_state.agents[i - 1], st.session_state.agents[i] = (
                    st.session_state.agents[i],
                    st.session_state.agents[i - 1],
                )
                st.rerun()

        with top[2]:
            if st.button(
                "⬇️",
                key=f"down_{agent_id}",
                disabled=i == len(st.session_state.agents) - 1,
                use_container_width=True,
            ):
                st.session_state.agents[i + 1], st.session_state.agents[i] = (
                    st.session_state.agents[i],
                    st.session_state.agents[i + 1],
                )
                st.rerun()

        with top[3]:
            if st.button("🗑️", key=f"delete_{agent_id}", use_container_width=True):
                st.session_state.agents.pop(i)
                st.session_state.rag_cache.pop(agent_id, None)
                st.rerun()

        left, right = st.columns([2, 1])

        with left:
            agent["system_prompt"] = st.text_area(
                "System Prompt",
                value=agent["system_prompt"],
                key=f"system_{agent_id}",
                height=150,
                help="이 Agent의 역할, 기준, 출력 형식 등을 지정하세요.",
            )

            agent["extra_prompt"] = st.text_area(
                "추가 Prompt (선택)",
                value=agent["extra_prompt"],
                key=f"extra_{agent_id}",
                height=110,
                help=(
                    "첫 Agent에는 최초 User Prompt와 함께, "
                    "두 번째 이후 Agent에는 바로 앞 Agent의 Output과 함께 전달됩니다."
                ),
            )

        with right:
            agent["model"] = st.text_input(
                "Model ID",
                value=agent["model"],
                key=f"model_{agent_id}",
                help="예: gpt-5-mini. 본인 API 계정에서 사용 가능한 모델 ID를 입력하세요.",
            )

            agent["max_output_tokens"] = st.number_input(
                "Max output tokens",
                min_value=128,
                max_value=32000,
                value=int(agent["max_output_tokens"]),
                step=128,
                key=f"max_tokens_{agent_id}",
            )

            if i == len(st.session_state.agents) - 1:
                agent["chatbot_mode"] = st.toggle(
                    "💬 Chatbot 모드",
                    value=bool(agent.get("chatbot_mode", False)),
                    key=f"chatbot_{agent_id}",
                    help=(
                        "마지막 Agent의 workflow 결과를 첫 응답으로 사용하고, "
                        "이전 workflow 결과 전체를 기억한 상태로 멀티턴 대화를 이어갑니다."
                    ),
                )
            else:
                agent["chatbot_mode"] = False

            agent["rag_enabled"] = st.toggle(
                "RAG 사용",
                value=bool(agent["rag_enabled"]),
                key=f"rag_{agent_id}",
            )

            if agent["rag_enabled"]:
                agent["top_k"] = st.slider(
                    "검색 Chunk 수 (Top-K)",
                    min_value=1,
                    max_value=8,
                    value=int(agent["top_k"]),
                    key=f"topk_{agent_id}",
                )

                uploaded_files = st.file_uploader(
                    "참조 파일 Drag & Drop",
                    type=["pdf", "docx", "pptx", "xlsx", "txt", "md", "csv", "json"],
                    accept_multiple_files=True,
                    key=f"files_{agent_id}",
                    help="파일은 해당 Agent의 RAG에만 사용됩니다.",
                )
                agent_uploads[agent_id] = uploaded_files
                if uploaded_files:
                    st.caption(f"{len(uploaded_files)}개 파일 선택됨")
            else:
                agent_uploads[agent_id] = []


st.divider()
st.markdown("### Workflow 실행")

initial_prompt = st.text_area(
    "최초 User Prompt",
    height=170,
    placeholder=(
        "예: 첨부 문서를 바탕으로 핵심 문제를 분석하고, "
        "최종적으로 실행 가능한 개선안을 만들어줘."
    ),
)

run_col, info_col = st.columns([1, 4])

with run_col:
    run_clicked = st.button(
        "▶️ Workflow 실행",
        type="primary",
        use_container_width=True,
        disabled=not st.session_state.agents,
    )

with info_col:
    st.caption(
        "실행 흐름: User Prompt → Agent 1 → Agent 2 → … → Final Output. "
        "각 단계의 출력이 다음 단계의 Workflow Input이 됩니다."
    )


if run_clicked:
    if not st.session_state.api_key.strip():
        st.error("먼저 왼쪽 사이드바에 OpenAI API Key를 입력해 주세요.")
        st.stop()

    empty_models = [a["name"] for a in st.session_state.agents if not a["model"].strip()]
    if empty_models:
        st.error(f"Model ID가 비어 있는 Agent가 있습니다: {', '.join(empty_models)}")
        st.stop()

    if not initial_prompt.strip():
        st.error("최초 User Prompt를 입력해 주세요.")
        st.stop()

    client = OpenAI(api_key=st.session_state.api_key.strip())
    current_input = initial_prompt.strip()
    results = []
    progress = st.progress(0, text="Workflow 시작")

    for i, agent in enumerate(st.session_state.agents):
        step_no = i + 1
        progress.progress(
            i / len(st.session_state.agents),
            text=f"{step_no}/{len(st.session_state.agents)} · {agent['name']} 실행 준비",
        )

        rag_context = ""
        retrieved = []
        rag_errors = []

        try:
            if agent["rag_enabled"]:
                files = agent_uploads.get(agent["id"], [])
                if not files:
                    st.warning(f"{agent['name']}: RAG가 켜져 있지만 참조 파일이 없어 RAG 없이 실행합니다.")
                else:
                    progress.progress(
                        i / len(st.session_state.agents),
                        text=f"{step_no}/{len(st.session_state.agents)} · {agent['name']} RAG 검색",
                    )
                    rag_query = current_input
                    if agent["extra_prompt"].strip():
                        rag_query += "\n\n" + agent["extra_prompt"].strip()

                    rag_context, retrieved, rag_errors = retrieve_context(
                        client=client,
                        agent_id=agent["id"],
                        files=files,
                        query=rag_query,
                        top_k=int(agent["top_k"]),
                    )

            progress.progress(
                i / len(st.session_state.agents),
                text=f"{step_no}/{len(st.session_state.agents)} · {agent['name']} LLM 호출",
            )

            output_text, actual_user_input = call_agent(
                client=client,
                agent=agent,
                workflow_input=current_input,
                rag_context=rag_context,
            )

            results.append(
                {
                    "step": step_no,
                    "name": agent["name"],
                    "model": agent["model"],
                    "workflow_input": current_input,
                    "actual_user_input": actual_user_input,
                    "output": output_text,
                    "retrieved": retrieved,
                    "rag_errors": rag_errors,
                }
            )
            current_input = output_text

        except Exception as exc:
            st.session_state.last_results = results
            progress.empty()
            st.error(f"{agent['name']} 실행 중 오류가 발생했습니다: {exc}")
            st.stop()

    st.session_state.last_results = results

    final_agent = st.session_state.agents[-1]
    if final_agent.get("chatbot_mode", False):
        memory_parts = [
            f"### Agent {result['step']} · {result['name']}\n{result['output']}"
            for result in results
        ]
        st.session_state.chat_context = "\n\n".join(memory_parts)
        st.session_state.chat_messages = []
        st.session_state.chat_agent_id = final_agent["id"]
    else:
        st.session_state.chat_context = ""
        st.session_state.chat_messages = []
        st.session_state.chat_agent_id = None

    progress.progress(1.0, text="Workflow 완료")
    st.success("Linear workflow 실행이 완료되었습니다.")


if st.session_state.last_results:
    st.markdown("### 실행 결과")

    for result in st.session_state.last_results:
        with st.expander(
            f"{result['step']}. {result['name']} · {result['model']}",
            expanded=result["step"] == len(st.session_state.last_results),
        ):
            input_tab, output_tab, rag_tab = st.tabs(["Input", "Output", "RAG"])

            with input_tab:
                st.text_area(
                    "이 Agent가 받은 Workflow Input",
                    value=result["workflow_input"],
                    height=180,
                    key=f"result_input_{result['step']}_{result['name']}",
                    disabled=True,
                )
                with st.expander("실제 API로 전달된 User Input 보기"):
                    st.code(result["actual_user_input"], language="text")

            with output_tab:
                st.markdown(result["output"])

            with rag_tab:
                if result["retrieved"]:
                    for ref in result["retrieved"]:
                        st.markdown(
                            f"**#{ref['rank']} · {ref['source']} · chunk {ref['chunk_no']} "
                            f"· similarity {ref['score']:.3f}**"
                        )
                        st.write(ref["text"])
                else:
                    st.caption("이 단계에서는 RAG 검색 결과가 없습니다.")

                if result["rag_errors"]:
                    st.warning("\n".join(result["rag_errors"]))

    final_output = st.session_state.last_results[-1]["output"]
    st.markdown("### Final Output")
    st.markdown(final_output)

    st.download_button(
        "⬇️ Final Output 다운로드",
        data=final_output,
        file_name="workflow_final_output.md",
        mime="text/markdown",
    )

    final_agent = st.session_state.agents[-1] if st.session_state.agents else None
    chatbot_active = (
        final_agent
        and final_agent.get("chatbot_mode", False)
        and st.session_state.chat_agent_id == final_agent["id"]
    )

    if chatbot_active:
        st.divider()
        st.markdown(f"### 💬 {final_agent['name']} Chatbot")
        st.caption(
            "완료된 workflow의 모든 Agent 결과를 세션 메모리로 유지하면서 후속 대화를 이어갑니다. "
            "브라우저 세션이 종료되면 이 대화 메모리는 초기화됩니다."
        )

        with st.chat_message("assistant"):
            st.markdown(final_output)

        for message in st.session_state.chat_messages:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

        chat_prompt = st.chat_input("워크플로우 결과에 대해 이어서 질문하세요.")

        if chat_prompt:
            if not st.session_state.api_key.strip():
                st.error("왼쪽 사이드바에 OpenAI API Key를 입력해 주세요.")
                st.stop()

            client = OpenAI(api_key=st.session_state.api_key.strip())
            chat_rag_context = ""

            try:
                if final_agent["rag_enabled"]:
                    files = agent_uploads.get(final_agent["id"], [])
                    if files:
                        chat_rag_context, _, _ = retrieve_context(
                            client=client,
                            agent_id=final_agent["id"],
                            files=files,
                            query=chat_prompt,
                            top_k=int(final_agent["top_k"]),
                        )

                assistant_reply = call_chat_agent(
                    client=client,
                    agent=final_agent,
                    workflow_memory=st.session_state.chat_context,
                    chat_messages=st.session_state.chat_messages,
                    user_prompt=chat_prompt,
                    rag_context=chat_rag_context,
                )

                st.session_state.chat_messages.append(
                    {"role": "user", "content": chat_prompt}
                )
                st.session_state.chat_messages.append(
                    {"role": "assistant", "content": assistant_reply}
                )
                st.rerun()

            except Exception as exc:
                st.error(f"Chatbot 응답 중 오류가 발생했습니다: {exc}")
